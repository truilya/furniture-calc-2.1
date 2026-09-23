"""
PDF-ТЗ → подготовленный XLSX через GPTunnel.

Установка:
    pip install streamlit pypdf openpyxl pandas requests

Запуск:
    streamlit run app.py

Адрес API GPTunnel можно изменить в боковой панели. По умолчанию
используется OpenAI-совместимый endpoint:
    https://gptunnel.ru/v1/chat/completions

Если в вашей конфигурации GPTunnel используется иной API URL или
другие идентификаторы моделей, укажите соответствующий URL и выберите
«Указать модель вручную…».

Важно: приложение не вставляет строки в Excel, не исправляет формулы
и не подменяет данные несоответствующего шаблона.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
from dataclasses import dataclass
from typing import Any
from urllib.parse import urlparse

import openpyxl
import pandas as pd
import requests
import streamlit as st
from openpyxl.cell.cell import MergedCell
from pypdf import PdfReader


# ---------------------------------------------------------------------
# Модели: список из предыдущей версии сохранён
# ---------------------------------------------------------------------

PROVIDERS: dict[str, dict[str, Any]] = {
    "OpenAI": {
        "models": [
            "gpt-4o-mini",
            "gpt-4o",
            "gpt-5.6-sol",
            "gpt-5.6-luna",
        ],
    },
    "OpenRouter": {
        "models": [
            "anthropic/claude-sonnet-5",
            "openai/gpt-5.6-sol",
            "google/gemini-3.1-flash-lite",
            "deepseek/deepseek-v4-pro",
            "qwen/qwen-2.5-vl-72b-instruct",
            "openai/gpt-4o-mini",
            "anthropic/claude-3.5-haiku",
            "meta-llama/llama-3.1-70b-instruct",
        ],
    },
    "Groq": {
        "models": [
            "llama-3.3-70b-versatile",
            "llama-3.1-8b-instant",
            "mixtral-8x7b-32768",
        ],
    },
    "Свой endpoint": {
        "models": [
            "deepseek-r1:70b",
            "qwen2.5-vl:72b",
        ],
    },
}

# Подписи групп сохранены для удобства поиска модели.
# Независимо от выбранной группы запрос всегда отправляется в GPTunnel.
DEFAULT_GPTUNNEL_CHAT_URL = (
    "https://gptunnel.ru/v1/chat/completions"
)

PRICE_SHEET = "Цены изделий (входящие)"
LOGISTICS_SHEET = "Логистика и пр. расходы"
ECONOMICS_SHEET = "Экономика"

MAX_PDF_BYTES = 25 * 1024 * 1024
MAX_XLSX_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_CHUNK_CHARS = 10_000
MAX_OUTPUT_TOKENS = 8_000
REQUEST_TIMEOUT = (10, 180)

ITEM_NUMBER_RE = re.compile(
    r"^\s*(\d{1,4})\s*[.)](?:\s|$)"
)


SYSTEM_PROMPT = """
Ты — интеллектуальный модуль парсинга государственных закупок (ТЗ) на поставку мебели.
Твоя задача — извлечь данные из текста ТЗ и вернуть строго валидный JSON:

{
  "procurement_info": {
    "object_of_purchase": "Поставка диванов и кресел",
    "customer": "ГБУЗ Поликлиника Троицкая ДЗМ",
    "total_quantity": 50
  },
  "items": [
    {
      "position_number": "1",
      "name_from_tz": "Секция диванная, деревянный каркас",
      "specifications": [
        "Вид товара: Секция диванная",
        "Тип каркаса: Деревянный",
        "Вид материала обивки: Кожа натуральная",
        "Ширина: Равно 2080 ММ",
        "Глубина: Равно 830 ММ",
        "Высота: Равно 770 ММ"
      ],
      "quantity": 1,
      "unit": "Штука",
      "delivery_logistics": {
        "address": "город Москва, улица 2-я Брестская, дом 6",
        "deadline_days_or_date": "c 1-го по 30-й рабочий день c даты заключения контракта",
        "delivery_form": "В разобранном виде"
      },
      "dimensions_mm": {
        "width_min": 2080,
        "width_max": 2080,
        "depth_min": 830,
        "depth_max": 830,
        "height_min": 770,
        "height_max": 770
      },
      "co_services": [
        "Сборка поставляемого товара в помещениях заказчика",
        "Расстановка поставляемого товара в помещениях заказчика",
        "Погрузочно-разгрузочные работы, включая подъем на этаж"
      ]
    }
  ]
}

Правила:
1. Каждая нумерованная товарная позиция приложения «Перечень объектов
   закупки» — отдельный объект. Не объединяй одинаковые названия.
2. Сохраняй номер позиции из документа. Не перенумеровывай товары.
3. Количество бери из строки объема товара. Если в ТЗ написано
   «50 (Штука)», верни "quantity": 50 и "unit": "Штука".
   quantity всегда должно быть JSON-числом, не строкой и не null.
   Не используй количество мест, ножек или дней.
4. Сохраняй индивидуальные характеристики позиции отдельными строками,
   включая материалы, габариты, диапазоны, цвет, отрицания и конструкцию.
5. Общие юридические разделы не включай в характеристики товара.
6. Адрес и срок бери из соответствующей позиции; если точного значения
   нет, верни null. Ничего не придумывай.
7. Цены, предполагаемую массу и параметры «по аналогии» не добавляй.
8. Если фрагмент содержит только продолжение позиции, используй ее
   действительный номер. Не добавляй позиции вне этого фрагмента.
9. Если в ТЗ написано «50 (Штука)», верни отдельные поля:
"quantity": 50,
"unit": "Штука".
Никогда не возвращай null для unit, если единица указана рядом с количеством.
""".strip()


class AppError(Exception):
    """Ошибка, предназначенная для показа в интерфейсе."""


@dataclass(frozen=True)
class GPTunnelSettings:
    chat_url: str
    api_key: str
    model: str


@dataclass(frozen=True)
class ProcurementInfo:
    object_of_purchase: str | None
    customer: str | None
    total_quantity: float | None


@dataclass(frozen=True)
class DeliveryLogistics:
    address: str | None
    deadline_days_or_date: str | None
    delivery_form: str | None


@dataclass(frozen=True)
class DimensionsMM:
    width_min: float | None
    width_max: float | None
    depth_min: float | None
    depth_max: float | None
    height_min: float | None
    height_max: float | None


@dataclass(frozen=True)
class ProcurementItem:
    position_number: int
    name_from_tz: str
    specifications: tuple[str, ...]
    quantity: float
    unit: str
    delivery_logistics: DeliveryLogistics
    dimensions_mm: DimensionsMM
    co_services: tuple[str, ...]


@dataclass(frozen=True)
class ExtractionResult:
    procurement_info: ProcurementInfo
    items: tuple[ProcurementItem, ...]


@dataclass(frozen=True)
class RowBlock:
    position_number: int
    first_row: int
    last_row: int


# ---------------------------------------------------------------------
# Файлы
# ---------------------------------------------------------------------

def read_uploaded_file(
    uploaded_file: Any,
    max_size: int,
    label: str,
) -> bytes:
    try:
        data = uploaded_file.getvalue()
    except Exception as exc:
        raise AppError(
            f"Не удалось прочитать {label}: {exc}"
        ) from exc

    if not data:
        raise AppError(f"{label} пуст.")

    if len(data) > max_size:
        raise AppError(
            f"{label} превышает размер "
            f"{max_size // (1024 * 1024)} МБ."
        )

    return data


def extract_pdf_text(pdf_bytes: bytes) -> str:
    if not pdf_bytes.startswith(b"%PDF"):
        raise AppError(
            "Загруженный файл не имеет сигнатуры PDF."
        )

    try:
        reader = PdfReader(
            io.BytesIO(pdf_bytes),
            strict=False,
        )

        if reader.is_encrypted:
            raise AppError(
                "Зашифрованный PDF не поддерживается."
            )

        if len(reader.pages) > MAX_PDF_PAGES:
            raise AppError(
                f"PDF содержит более {MAX_PDF_PAGES} страниц."
            )

        pages: list[str] = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            page_text = (
                page.extract_text(
                    extraction_mode="layout"
                )
                or ""
            )

            if page_text.strip():
                pages.append(
                    f"\n[СТРАНИЦА {page_number}]\n"
                    f"{page_text}"
                )

        text = "\n".join(pages).strip()

    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            f"Ошибка чтения PDF: {exc}"
        ) from exc

    if len(text) < 100:
        raise AppError(
            "В PDF недостаточно извлекаемого текста. "
            "Для скана сначала выполните OCR."
        )

    return text


def open_template(
    xlsx_bytes: bytes,
) -> openpyxl.Workbook:
    if not xlsx_bytes.startswith(b"PK"):
        raise AppError(
            "Файл шаблона не похож на XLSX."
        )

    try:
        return openpyxl.load_workbook(
            io.BytesIO(xlsx_bytes),
            data_only=False,
            read_only=False,
            keep_links=True,
        )
    except Exception as exc:
        raise AppError(
            f"Не удалось открыть XLSX-шаблон: {exc}"
        ) from exc


# ---------------------------------------------------------------------
# GPTunnel
# ---------------------------------------------------------------------

def validate_gptunnel_url(
    raw_url: str,
) -> str:
    url = raw_url.strip().rstrip("/")
    parsed = urlparse(url)

    if (
        parsed.scheme != "https"
        or not parsed.netloc
    ):
        raise AppError(
            "Укажите HTTPS URL метода GPTunnel "
            "chat/completions."
        )

    if parsed.username or parsed.password:
        raise AppError(
            "Не передавайте API-ключ внутри URL."
        )

    if (
        parsed.hostname is None
        or not (
            parsed.hostname == "gptunnel.ru"
            or parsed.hostname.endswith(
                ".gptunnel.ru"
            )
        )
    ):
        raise AppError(
            "Запросы должны отправляться только "
            "на домен GPTunnel."
        )

    if not parsed.path.endswith(
        "/chat/completions"
    ):
        raise AppError(
            "URL должен оканчиваться "
            "на /chat/completions."
        )

    return url


def parse_json_response(
    content: str,
) -> dict[str, Any]:
    text = content.strip()

    fenced = re.fullmatch(
        r"```(?:json)?\s*(\{.*\})\s*```",
        text,
        flags=(
            re.DOTALL
            | re.IGNORECASE
        ),
    )

    if fenced:
        text = fenced.group(1)

    try:
        result = json.loads(text)
    except json.JSONDecodeError as exc:
        raise AppError(
            "Модель вернула невалидный JSON: "
            f"{exc}"
        ) from exc

    if not isinstance(result, dict):
        raise AppError(
            "Ответ модели должен быть "
            "JSON-объектом."
        )

    return result


def call_gptunnel(
    settings: GPTunnelSettings,
    user_prompt: str,
) -> dict[str, Any]:
    headers = {
        "Authorization": (
            f"Bearer {settings.api_key}"
        ),
        "Content-Type": (
            "application/json"
        ),
    }

    payload = {
        "model": settings.model,
        "messages": [
            {
                "role": "system",
                "content": SYSTEM_PROMPT,
            },
            {
                "role": "user",
                "content": user_prompt,
            },
        ],
        "temperature": 0,
        "max_tokens": MAX_OUTPUT_TOKENS,
    }

    try:
        response = requests.post(
            settings.chat_url,
            headers=headers,
            json=payload,
            timeout=REQUEST_TIMEOUT,
        )

        response.raise_for_status()

        body = response.json()
        content = body["choices"][0][
            "message"
        ]["content"]

        if isinstance(content, list):
            content = "".join(
                part.get("text", "")
                for part in content
                if isinstance(
                    part,
                    dict,
                )
            )

        if (
            not isinstance(
                content,
                str,
            )
            or not content.strip()
        ):
            raise AppError(
                "GPTunnel вернул пустой ответ модели."
            )

        return parse_json_response(
            content
        )

    except requests.Timeout as exc:
        raise AppError(
            "Превышено время ожидания "
            "ответа GPTunnel."
        ) from exc

    except requests.HTTPError as exc:
        status = (
            exc.response.status_code
            if exc.response is not None
            else "неизвестен"
        )

        raise AppError(
            f"GPTunnel вернул HTTP {status} "
            f"для модели «{settings.model}». "
            "Проверьте ключ, URL, баланс и "
            "доступность ID модели."
        ) from exc

    except requests.RequestException as exc:
        raise AppError(
            f"Сетевая ошибка GPTunnel: {exc}"
        ) from exc

    except (
        KeyError,
        IndexError,
        TypeError,
        ValueError,
    ) as exc:
        raise AppError(
            "Неожиданный формат ответа "
            f"GPTunnel: {exc}"
        ) from exc


def split_text(
    text: str,
    max_chars: int = MAX_CHUNK_CHARS,
) -> list[str]:
    if len(text) <= max_chars:
        return [text]

    paragraphs = re.split(
        r"\n\s*\n",
        text,
    )

    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        paragraph = paragraph.strip()

        if not paragraph:
            continue

        if len(paragraph) > max_chars:
            if current:
                chunks.append(current)
                current = ""

            step = max_chars - 500

            for start in range(
                0,
                len(paragraph),
                step,
            ):
                chunks.append(
                    paragraph[
                        start:start + max_chars
                    ]
                )

            continue

        candidate = (
            f"{current}\n\n{paragraph}"
            if current
            else paragraph
        )

        if (
            len(candidate) > max_chars
            and current
        ):
            chunks.append(current)

            current = (
                "[КОНТЕКСТ ПРЕДЫДУЩЕГО "
                "ФРАГМЕНТА]\n"
                f"{current[-500:]}\n\n"
                f"{paragraph}"
            )
        else:
            current = candidate

    if current:
        chunks.append(current)

    return chunks


# ---------------------------------------------------------------------
# Проверка JSON
# ---------------------------------------------------------------------

def clean_string(
    value: Any,
    field: str,
    *,
    nullable: bool = False,
) -> str | None:
    if value is None and nullable:
        return None

    if not isinstance(value, str):
        raise AppError(
            f"Поле «{field}» должно быть строкой."
            f"value «{value}»"
        )

    text = re.sub(
        r"\s+",
        " ",
        value,
    ).strip()

    if not text and not nullable:
        raise AppError(
            f"Поле «{field}» пусто."
        )

    return text or None


def optional_nonnegative_number(
    value: Any,
    field: str,
) -> float | None:
    if value is None:
        return None

    if (
        isinstance(
            value,
            bool,
        )
        or not isinstance(
            value,
            (int, float),
        )
    ):
        raise AppError(
            f"Поле «{field}» должно быть числом "
            "или null."
        )

    number = float(value)

    if (
        not math.isfinite(number)
        or number < 0
    ):
        raise AppError(
            f"Поле «{field}» содержит "
            "недопустимое число."
        )

    return number


def validate_dimensions(
    raw: Any,
    position: int,
) -> DimensionsMM:
    if not isinstance(raw, dict):
        raise AppError(
            f"Позиция №{position}: "
            "dimensions_mm должен быть объектом."
        )

    fields = {}

    for name in (
        "width_min",
        "width_max",
        "depth_min",
        "depth_max",
        "height_min",
        "height_max",
    ):
        fields[name] = (
            optional_nonnegative_number(
                raw.get(name),
                (
                    f"позиция №{position}, "
                    f"dimensions_mm.{name}"
                ),
            )
        )

    for axis in (
        "width",
        "depth",
        "height",
    ):
        minimum = fields[
            f"{axis}_min"
        ]

        maximum = fields[
            f"{axis}_max"
        ]

        if (
            minimum is not None
            and maximum is not None
            and minimum > maximum
        ):
            raise AppError(
                f"Позиция №{position}: "
                f"{axis}_min больше "
                f"{axis}_max."
            )

    return DimensionsMM(
        **fields,
    )


def validate_string_list(
    raw: Any,
    field: str,
    *,
    allow_empty: bool,
) -> tuple[str, ...]:
    if not isinstance(
        raw,
        list,
    ):
        raise AppError(
            f"Поле «{field}» должно быть массивом."
        )

    if (
        not allow_empty
        and not raw
    ):
        raise AppError(
            f"Поле «{field}» пусто."
        )

    result: list[str] = []
    seen: set[str] = set()

    for value in raw:
        line = clean_string(
            value,
            field,
        )

        assert line is not None

        normalized = (
            line.casefold()
        )

        if normalized not in seen:
            seen.add(normalized)
            result.append(line)

    return tuple(result)

def parse_quantity_and_unit(
    quantity_value: Any,
    unit_value: Any,
    position: int,
) -> tuple[float, str]:
    """Разбирает JSON-число либо строку вида '50 (Штука)'."""

    embedded_unit: str | None = None

    if isinstance(quantity_value, bool) or quantity_value is None:
        raise AppError(
            f"Позиция №{position}: quantity не указан или имеет неверный тип."
        )

    if isinstance(quantity_value, (int, float)):
        quantity = float(quantity_value)

    elif isinstance(quantity_value, str):
        text = quantity_value.strip().replace("\u00a0", " ")

        match = re.fullmatch(
            r"\s*(\d+(?:[ \u202f]\d{3})*(?:[.,]\d+)?)"
            r"\s*(?:\(([^)]+)\))?\s*",
            text,
        )

        if not match:
            raise AppError(
                f"Позиция №{position}: не удалось распознать quantity: "
                f"{text[:80]!r}"
            )

        quantity = float(
            match.group(1)
            .replace(" ", "")
            .replace("\u202f", "")
            .replace(",", ".")
        )
        embedded_unit = match.group(2)

    else:
        raise AppError(
            f"Позиция №{position}: неверный тип quantity: "
            f"{type(quantity_value).__name__}."
        )

    if not math.isfinite(quantity) or not 0 < quantity < 1_000_000_000:
        raise AppError(
            f"Позиция №{position}: количество должно быть положительным числом."
        )

    if isinstance(unit_value, str) and unit_value.strip():
        unit = unit_value.strip()
    elif embedded_unit:
        unit = embedded_unit.strip()
    else:
        raise AppError(
            f"Позиция №{position}: единица измерения отсутствует. "
            "В ответе модели ожидается поле unit, например «Штука»."
        )

    return quantity, unit

def validate_item(
    raw: Any,
) -> ProcurementItem:
    if not isinstance(raw, dict):
        raise AppError(
            "Одна из позиций не является "
            "JSON-объектом."
        )

    raw_number = raw.get(
        "position_number"
    )

    if isinstance(
        raw_number,
        bool,
    ):
        raise AppError(
            "Некорректный номер позиции."
        )

    if isinstance(
        raw_number,
        int,
    ):
        number = raw_number

    elif (
        isinstance(
            raw_number,
            str,
        )
        and raw_number.strip().isdigit()
    ):
        number = int(
            raw_number.strip()
        )

    else:
        raise AppError(
            "position_number должен содержать "
            "номер позиции."
        )

    if number < 1:
        raise AppError(
            "Номер позиции должен быть "
            "положительным."
        )

    logistics_raw = raw.get(
        "delivery_logistics"
    )

    if not isinstance(
        logistics_raw,
        dict,
    ):
        raise AppError(
            f"Позиция №{number}: "
            "нет объекта delivery_logistics."
        )

    name = clean_string(
        raw.get(
            "name_from_tz"
        ),
        (
            f"позиция №{number}, "
            "name_from_tz"
        ),
    )

    st.write(
        f"Диагностика позиции №{number}:",
        {
            "quantity": raw.get("quantity"),
            "unit": raw.get("unit"),
        },
    )

    quantity, unit = parse_quantity_and_unit(
        raw.get("quantity"),
        raw.get("unit"),
        number,
    )

    assert name is not None
    assert unit is not None

    return ProcurementItem(
        position_number=number,
        name_from_tz=name,
        specifications=(
            validate_string_list(
                raw.get(
                    "specifications"
                ),
                (
                    f"позиция №{number}, "
                    "specifications"
                ),
                allow_empty=False,
            )
        ),
        quantity=quantity,
        unit=unit,
        delivery_logistics=(
            DeliveryLogistics(
                address=clean_string(
                    logistics_raw.get(
                        "address"
                    ),
                    (
                        f"позиция №{number}, "
                        "address"
                    ),
                    nullable=True,
                ),
                deadline_days_or_date=(
                    clean_string(
                        logistics_raw.get(
                            "deadline_days_or_date"
                        ),
                        (
                            f"позиция №{number}, "
                            "deadline_days_or_date"
                        ),
                        nullable=True,
                    )
                ),
                delivery_form=clean_string(
                    logistics_raw.get(
                        "delivery_form"
                    ),
                    (
                        f"позиция №{number}, "
                        "delivery_form"
                    ),
                    nullable=True,
                ),
            )
        ),
        dimensions_mm=(
            validate_dimensions(
                raw.get(
                    "dimensions_mm"
                ),
                number,
            )
        ),
        co_services=(
            validate_string_list(
                raw.get(
                    "co_services"
                ),
                (
                    f"позиция №{number}, "
                    "co_services"
                ),
                allow_empty=True,
            )
        ),
    )


def validate_procurement_info(
    raw: Any,
) -> ProcurementInfo:
    if not isinstance(
        raw,
        dict,
    ):
        raise AppError(
            "Ответ модели не содержит "
            "procurement_info."
        )

    return ProcurementInfo(
        object_of_purchase=(
            clean_string(
                raw.get(
                    "object_of_purchase"
                ),
                "object_of_purchase",
                nullable=True,
            )
        ),
        customer=clean_string(
            raw.get(
                "customer"
            ),
            "customer",
            nullable=True,
        ),
        total_quantity=(
            optional_nonnegative_number(
                raw.get(
                    "total_quantity"
                ),
                "total_quantity",
            )
        ),
    )


def ordered_union(
    first: tuple[str, ...],
    second: tuple[str, ...],
) -> tuple[str, ...]:
    result = list(first)

    seen = {
        value.casefold()
        for value in result
    }

    for value in second:
        normalized = (
            value.casefold()
        )

        if normalized not in seen:
            result.append(value)
            seen.add(normalized)

    return tuple(result)


def merge_items(
    old: ProcurementItem,
    new: ProcurementItem,
) -> ProcurementItem:
    if (
        old.name_from_tz.casefold()
        != new.name_from_tz.casefold()
        or old.quantity
        != new.quantity
        or old.unit.casefold()
        != new.unit.casefold()
    ):
        raise AppError(
            "Противоречивые ответы модели "
            f"для позиции №{old.position_number}."
        )

    old_logistics = (
        old.delivery_logistics
    )

    new_logistics = (
        new.delivery_logistics
    )

    # Для пересечения фрагментов не выбираем молча
    # два различных конкретных значения.
    for field in (
        "address",
        "deadline_days_or_date",
        "delivery_form",
    ):
        left = getattr(
            old_logistics,
            field,
        )

        right = getattr(
            new_logistics,
            field,
        )

        if (
            left
            and right
            and left.casefold()
            != right.casefold()
        ):
            raise AppError(
                f"Позиция №{old.position_number}: "
                f"противоречие в поле {field}."
            )

    old_dims = (
        old.dimensions_mm
    )

    new_dims = (
        new.dimensions_mm
    )

    dimension_values: dict[
        str,
        float | None
    ] = {}

    for field in (
        "width_min",
        "width_max",
        "depth_min",
        "depth_max",
        "height_min",
        "height_max",
    ):
        left = getattr(
            old_dims,
            field,
        )

        right = getattr(
            new_dims,
            field,
        )

        if (
            left is not None
            and right is not None
            and left != right
        ):
            raise AppError(
                f"Позиция №{old.position_number}: "
                f"противоречие в {field}."
            )

        dimension_values[field] = (
            left
            if left is not None
            else right
        )

    return ProcurementItem(
        position_number=(
            old.position_number
        ),
        name_from_tz=(
            old.name_from_tz
        ),
        specifications=(
            ordered_union(
                old.specifications,
                new.specifications,
            )
        ),
        quantity=old.quantity,
        unit=old.unit,
        delivery_logistics=(
            DeliveryLogistics(
                address=(
                    old_logistics.address
                    or new_logistics.address
                ),
                deadline_days_or_date=(
                    old_logistics.deadline_days_or_date
                    or new_logistics.deadline_days_or_date
                ),
                delivery_form=(
                    old_logistics.delivery_form
                    or new_logistics.delivery_form
                ),
            )
        ),
        dimensions_mm=(
            DimensionsMM(
                **dimension_values
            )
        ),
        co_services=(
            ordered_union(
                old.co_services,
                new.co_services,
            )
        ),
    )


def extract_data_with_llm(
    pdf_text: str,
    settings: GPTunnelSettings,
) -> ExtractionResult:
    chunks = split_text(
        pdf_text
    )

    merged_items: dict[
        int,
        ProcurementItem
    ] = {}

    procurement_infos: list[
        ProcurementInfo
    ] = []

    for index, chunk in enumerate(
        chunks,
        start=1,
    ):
        raw_result = call_gptunnel(
            settings,
            (
                "Извлеки данные только из "
                f"фрагмента {index}/{len(chunks)}.\n"
                "Если общий объект закупки или "
                "заказчик отсутствует во фрагменте, "
                "используй null. "
                "Если товарных позиций нет, "
                "верни items: [].\n\n"
                f"{chunk}"
            ),
        )

        info = (
            validate_procurement_info(
                raw_result.get(
                    "procurement_info"
                )
            )
        )

        procurement_infos.append(
            info
        )

        raw_items = (
            raw_result.get(
                "items"
            )
        )

        if not isinstance(
            raw_items,
            list,
        ):
            raise AppError(
                f"Фрагмент {index}: поле items "
                "должно быть массивом."
            )

        for raw_item in raw_items:

            st.write(
                f"Диагностика позиции №{number}:",
                {
                    "quantity": raw.get("quantity"),
                    "unit": raw.get("unit"),
                },
            )

            item = validate_item(
                raw_item
            )

            existing = (
                merged_items.get(
                    item.position_number
                )
            )

            if existing is None:
                merged_items[
                    item.position_number
                ] = item
            else:
                merged_items[
                    item.position_number
                ] = merge_items(
                    existing,
                    item,
                )

    if not merged_items:
        raise AppError(
            "Модель не обнаружила "
            "товарных позиций."
        )

    items = tuple(
        merged_items[number]
        for number in sorted(
            merged_items
        )
    )

    def first_nonempty(
        field: str,
    ) -> str | None:
        for info in procurement_infos:
            value = getattr(
                info,
                field,
            )

            if value:
                return value

        return None

    # Сумма проверенных количеств надежнее, чем
    # total_quantity отдельного фрагмента.
    total_quantity = sum(
        item.quantity
        for item in items
    )

    return ExtractionResult(
        procurement_info=(
            ProcurementInfo(
                object_of_purchase=(
                    first_nonempty(
                        "object_of_purchase"
                    )
                ),
                customer=(
                    first_nonempty(
                        "customer"
                    )
                ),
                total_quantity=(
                    total_quantity
                ),
            )
        ),
        items=items,
    )


# ---------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------

def item_number(
    value: Any,
) -> int | None:
    if not isinstance(
        value,
        str,
    ):
        return None

    match = ITEM_NUMBER_RE.match(
        value
    )

    return (
        int(match.group(1))
        if match
        else None
    )


def find_header_row(
    ws: Any,
) -> int:
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(
            ws.max_row,
            30,
        ),
        max_col=min(
            ws.max_column,
            7,
        ),
    ):
        values = [
            str(
                cell.value
                or ""
            )
            .replace(
                "\n",
                " ",
            )
            .casefold()
            for cell in row
        ]

        if any(
            (
                "наименование товара "
                "из тз"
            )
            in value
            for value in values
        ):
            return row[0].row

    raise AppError(
        f"Лист «{ws.title}»: "
        "не найден заголовок "
        "«Наименование товара из ТЗ»."
    )


def detect_blocks(
    ws: Any,
    header_row: int,
) -> list[RowBlock]:
    starts: list[
        tuple[int, int]
    ] = []

    stop_row = (
        ws.max_row + 1
    )

    for row in range(
        header_row + 1,
        ws.max_row + 1,
    ):
        first_value = (
            ws.cell(
                row,
                1,
            ).value
        )

        number = item_number(
            first_value
        )

        if number is not None:
            starts.append(
                (
                    row,
                    number,
                )
            )
            continue

        if (
            starts
            and isinstance(
                first_value,
                str,
            )
        ):
            lowered = (
                first_value
                .strip()
                .casefold()
            )

            if lowered.startswith(
                (
                    "итого",
                    "итоговая",
                    "всего",
                )
            ):
                stop_row = row
                break

    if not starts:
        raise AppError(
            f"Лист «{ws.title}»: "
            "нет нумерованных блоков."
        )

    numbers = [
        number
        for _,
        number in starts
    ]

    if len(
        numbers
    ) != len(
        set(numbers)
    ):
        raise AppError(
            f"Лист «{ws.title}»: "
            "номера блоков повторяются."
        )

    result: list[
        RowBlock
    ] = []

    for index, (
        row,
        number,
    ) in enumerate(starts):
        last_row = (
            starts[
                index + 1
            ][0] - 1
            if index + 1
            < len(starts)
            else stop_row - 1
        )

        result.append(
            RowBlock(
                position_number=(
                    number
                ),
                first_row=row,
                last_row=(
                    last_row
                ),
            )
        )

    return result


def assert_editable(
    ws: Any,
    row: int,
    column: int,
) -> None:
    cell = ws.cell(
        row,
        column,
    )

    if isinstance(
        cell,
        MergedCell,
    ):
        raise AppError(
            f"Ячейка "
            f"{ws.title}!"
            f"{cell.coordinate} "
            "входит в объединенный "
            "диапазон."
        )

    if (
        cell.data_type
        == "f"
        or (
            isinstance(
                cell.value,
                str,
            )
            and cell.value
            .startswith("=")
        )
    ):
        raise AppError(
            f"Ячейка "
            f"{ws.title}!"
            f"{cell.coordinate} "
            "содержит формулу."
        )


def set_cell(
    ws: Any,
    row: int,
    column: int,
    value: (
        str
        | int
        | float
        | None
    ),
) -> None:
    assert_editable(
        ws,
        row,
        column,
    )

    ws.cell(
        row,
        column,
    ).value = value


def write_lines(
    ws: Any,
    block: RowBlock,
    column: int,
    values: list[str],
) -> None:
    available = (
        block.last_row
        - block.first_row
        + 1
    )

    if len(
        values
    ) > available:
        raise AppError(
            f"Лист «{ws.title}», "
            f"позиция №"
            f"{block.position_number}: "
            f"нужно {len(values)} "
            f"строк, доступно "
            f"{available}. "
            "Расширьте шаблон вручную."
        )

    for row in range(
        block.first_row,
        block.last_row + 1,
    ):
        assert_editable(
            ws,
            row,
            column,
        )

    for index, value in enumerate(
        values
    ):
        set_cell(
            ws,
            block.first_row
            + index,
            column,
            value,
        )

    for row in range(
        block.first_row
        + len(values),
        block.last_row
        + 1,
    ):
        set_cell(
            ws,
            row,
            column,
            None,
        )


def dimension_lines(
    specifications: tuple[
        str,
        ...
    ],
) -> list[str]:
    prefixes = (
        "высота",
        "глубина",
        "ширина",
        "длина",
        "толщина",
        "диаметр",
    )

    return [
        line
        for line in specifications
        if line.casefold()
        .startswith(prefixes)
    ]


def validate_mapping(
    workbook: openpyxl.Workbook,
    items: tuple[
        ProcurementItem,
        ...
    ],
) -> dict[
    str,
    dict[
        int,
        RowBlock,
    ],
]:
    required = {
        PRICE_SHEET,
        LOGISTICS_SHEET,
        ECONOMICS_SHEET,
    }

    missing = required - set(
        workbook.sheetnames
    )

    if missing:
        raise AppError(
            "В шаблоне нет листов: "
            + ", ".join(
                sorted(missing)
            )
        )

    mappings: dict[
        str,
        dict[
            int,
            RowBlock,
        ],
    ] = {}

    for sheet_name in (
        PRICE_SHEET,
        LOGISTICS_SHEET,
        ECONOMICS_SHEET,
    ):
        ws = workbook[
            sheet_name
        ]

        blocks = detect_blocks(
            ws,
            find_header_row(
                ws
            ),
        )

        mappings[
            sheet_name
        ] = {
            block.position_number: block
            for block in blocks
        }

    expected = {
        item.position_number
        for item in items
    }

    for sheet_name, mapping in (
        mappings.items()
    ):
        actual = set(
            mapping
        )

        if actual != expected:
            raise AppError(
                f"Лист «{sheet_name}» "
                "не соответствует "
                "позициям PDF. "
                "Нет блоков: "
                f"{sorted(expected - actual)}; "
                "лишние блоки: "
                f"{sorted(actual - expected)}. "
                "Подмена данных "
                "другой закупки "
                "не выполняется."
            )

    return mappings


def populate_workbook(
    workbook: openpyxl.Workbook,
    extraction: ExtractionResult,
) -> None:
    mapping = validate_mapping(
        workbook,
        extraction.items,
    )

    price_ws = workbook[
        PRICE_SHEET
    ]

    logistics_ws = workbook[
        LOGISTICS_SHEET
    ]

    economics_ws = workbook[
        ECONOMICS_SHEET
    ]

    for item in (
        extraction.items
    ):
        price_block = (
            mapping[
                PRICE_SHEET
            ][
                item.position_number
            ]
        )

        logistics_block = (
            mapping[
                LOGISTICS_SHEET
            ][
                item.position_number
            ]
        )

        economics_block = (
            mapping[
                ECONOMICS_SHEET
            ][
                item.position_number
            ]
        )

        label = (
            f"{item.position_number}.\n"
            f"{item.name_from_tz}"
        )

        quantity: (
            int
            | float
        ) = (
            int(
                item.quantity
            )
            if item.quantity
            .is_integer()
            else item.quantity
        )

        # Цены изделий:
        # A — позиция, D — характеристики,
        # E — количество.
        set_cell(
            price_ws,
            price_block.first_row,
            1,
            label,
        )

        set_cell(
            price_ws,
            price_block.first_row,
            5,
            quantity,
        )

        write_lines(
            price_ws,
            price_block,
            4,
            list(
                item.specifications
            ),
        )

        # Логистика:
        # A — позиция, C — строки габаритов,
        # D — количество, E — единица.
        set_cell(
            logistics_ws,
            logistics_block.first_row,
            1,
            label,
        )

        set_cell(
            logistics_ws,
            logistics_block.first_row,
            4,
            quantity,
        )

        set_cell(
            logistics_ws,
            logistics_block.first_row,
            5,
            item.unit,
        )

        write_lines(
            logistics_ws,
            logistics_block,
            3,
            dimension_lines(
                item.specifications
            ),
        )

        # Экономика:
        # A — позиция, C — количество.
        # Цены и формулы не меняем.
        set_cell(
            economics_ws,
            economics_block.first_row,
            1,
            label,
        )

        set_cell(
            economics_ws,
            economics_block.first_row,
            3,
            quantity,
        )

    try:
        workbook.calculation = (
            openpyxl.workbook
            .properties
            .CalcProperties(
                calcMode="auto",
                fullCalcOnLoad=True,
                forceFullCalc=True,
            )
        )
    except AttributeError:
        pass


def workbook_to_bytes(
    workbook: openpyxl.Workbook,
) -> bytes:
    try:
        output = io.BytesIO()

        workbook.save(
            output
        )

        result = (
            output.getvalue()
        )

        if not result:
            raise ValueError(
                "Пустой XLSX."
            )

        return result

    except Exception as exc:
        raise AppError(
            "Ошибка сохранения XLSX: "
            f"{exc}"
        ) from exc


def items_dataframe(
    extraction: ExtractionResult,
) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "№": (
                    item.position_number
                ),
                "Наименование": (
                    item.name_from_tz
                ),
                "Количество": (
                    item.quantity
                ),
                "Ед. изм.": (
                    item.unit
                ),
                "Адрес": (
                    item.delivery_logistics
                    .address
                    or "Не указан"
                ),
                "Срок": (
                    item.delivery_logistics
                    .deadline_days_or_date
                    or "Не указан"
                ),
                "Характеристик": len(
                    item.specifications
                ),
                "Сопутствующих услуг": len(
                    item.co_services
                ),
            }
            for item in (
                extraction.items
            )
        ]
    )


# ---------------------------------------------------------------------
# Streamlit
# ---------------------------------------------------------------------

def settings_ui() -> GPTunnelSettings:
    st.sidebar.header(
        "GPTunnel"
    )

    st.sidebar.caption(
        "Группы ниже используются только "
        "для выбора модели. Все запросы "
        "отправляются через GPTunnel."
    )

    model_group = (
        st.sidebar.selectbox(
            "Группа моделей",
            list(PROVIDERS),
        )
    )

    model_choice = (
        st.sidebar.selectbox(
            "Модель",
            [
                *PROVIDERS[
                    model_group
                ]["models"],
                (
                    "Указать модель "
                    "вручную…"
                ),
            ],
            key=(
                "model_choice_"
                f"{model_group}"
            ),
        )
    )

    if model_choice == (
        "Указать модель вручную…"
    ):
        model = (
            st.sidebar
            .text_input(
                "ID модели в GPTunnel",
                key=(
                    "manual_model_"
                    f"{model_group}"
                ),
            )
            .strip()
        )
    else:
        model = model_choice

    chat_url = (
        st.sidebar.text_input(
            "URL GPTunnel API",
            value=(
                DEFAULT_GPTUNNEL_CHAT_URL
            ),
            help=(
                "Полный HTTPS URL метода "
                "/chat/completions."
            ),
        )
    )

    api_key = (
        st.sidebar.text_input(
            "API-ключ GPTunnel",
            type="password",
            help=(
                "Ключ используется для "
                "запросов и не записывается "
                "в Excel."
            ),
        )
        .strip()
    )

    if not model:
        raise AppError(
            "Выберите или укажите "
            "модель."
        )

    if not api_key:
        raise AppError(
            "Укажите API-ключ "
            "GPTunnel."
        )

    return GPTunnelSettings(
        chat_url=(
            validate_gptunnel_url(
                chat_url
            )
        ),
        api_key=api_key,
        model=model,
    )


def main() -> None:
    st.set_page_config(
        page_title=(
            "ТЗ → Excel"
        ),
        page_icon="📋",
        layout="wide",
    )

    st.title(
        "Техническое задание "
        "→ Excel"
    )

    st.write(
        "Загрузите PDF с текстовым слоем "
        "и подготовленный XLSX-шаблон. "
        "Анализ выполняется через GPTunnel."
    )

    try:
        settings = (
            settings_ui()
        )
    except AppError as exc:
        st.sidebar.error(
            str(exc)
        )
        return

    st.subheader(
        "Исходные файлы"
    )

    pdf_upload = (
        st.file_uploader(
            "PDF с техническим заданием",
            type=["pdf"],
            key="pdf_upload",
        )
    )

    xlsx_upload = (
        st.file_uploader(
            "XLSX-шаблон",
            type=["xlsx"],
            key="xlsx_upload",
        )
    )

    if (
        pdf_upload is None
        or xlsx_upload is None
    ):
        return

    try:
        pdf_bytes = (
            read_uploaded_file(
                pdf_upload,
                MAX_PDF_BYTES,
                "PDF-файл",
            )
        )

        xlsx_bytes = (
            read_uploaded_file(
                xlsx_upload,
                MAX_XLSX_BYTES,
                "Excel-шаблон",
            )
        )

    except AppError as exc:
        st.error(
            str(exc)
        )
        return

    # Ключ участвует только в вычислении хеша
    # текущего запуска; его открытое значение
    # не сохраняется в результате или XLSX.
    fingerprint = (
        hashlib.sha256(
            b"\0".join(
                [
                    pdf_bytes,
                    xlsx_bytes,
                    settings.chat_url
                    .encode(),
                    settings.model
                    .encode(),
                    settings.api_key
                    .encode(),
                ]
            )
        )
        .hexdigest()
    )

    if (
        st.session_state.get(
            "result_fingerprint"
        )
        != fingerprint
    ):
        for key in (
            "result_binary",
            "result_dataframe",
            "result_info",
            "result_fingerprint",
        ):
            st.session_state.pop(
                key,
                None,
            )

    has_result = bool(
        st.session_state.get(
            "result_binary"
        )
    )

    if st.button(
        "Проанализировать ТЗ "
        "и заполнить шаблон",
        type="primary",
        disabled=has_result,
    ):
        with st.spinner(
            "Читаем PDF, отправляем "
            "фрагменты в GPTunnel "
            "и проверяем XLSX…"
        ):
            try:
                pdf_text = (
                    extract_pdf_text(
                        pdf_bytes
                    )
                )

                extraction = (
                    extract_data_with_llm(
                        pdf_text,
                        settings,
                    )
                )

                workbook = (
                    open_template(
                        xlsx_bytes
                    )
                )

                populate_workbook(
                    workbook,
                    extraction,
                )

                result_binary = (
                    workbook_to_bytes(
                        workbook
                    )
                )

                st.session_state[
                    "result_binary"
                ] = result_binary

                st.session_state[
                    "result_dataframe"
                ] = items_dataframe(
                    extraction
                )

                st.session_state[
                    "result_info"
                ] = (
                    extraction
                    .procurement_info
                )

                st.session_state[
                    "result_fingerprint"
                ] = fingerprint

            except AppError as exc:
                st.error(
                    str(exc)
                )

            except Exception as exc:
                st.error(
                    "Непредвиденная ошибка: "
                    f"{type(exc).__name__}: "
                    f"{exc}"
                )

    if st.session_state.get(
        "result_binary"
    ):
        st.success(
            "Файл подготовлен."
        )

        info: ProcurementInfo = (
            st.session_state[
                "result_info"
            ]
        )

        st.subheader(
            "Сведения о закупке"
        )

        st.write(
            {
                "Объект закупки": (
                    info.object_of_purchase
                ),
                "Заказчик": (
                    info.customer
                ),
                "Общее количество": (
                    info.total_quantity
                ),
            }
        )

        st.subheader(
            "Извлеченные позиции"
        )

        st.dataframe(
            st.session_state[
                "result_dataframe"
            ],
            use_container_width=True,
            hide_index=True,
        )

        st.warning(
            "Проверьте цены, логистические "
            "оценки и служебные листы: они "
            "не извлекаются из нового PDF и "
            "могут относиться к другой закупке. "
            "Формулы пересчитываются Excel "
            "при открытии файла."
        )

        st.download_button(
            label=(
                "Скачать заполненный XLSX"
            ),
            data=(
                st.session_state[
                    "result_binary"
                ]
            ),
            file_name=(
                "ТЗ_заполненный_"
                "шаблон.xlsx"
            ),
            mime=(
                "application/vnd."
                "openxmlformats-"
                "officedocument."
                "spreadsheetml.sheet"
            ),
        )


if __name__ == "__main__":
    main()